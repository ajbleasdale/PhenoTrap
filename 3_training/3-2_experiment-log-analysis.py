from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# =========================================================
# SETTINGS
# =========================================================
FILE_PATH = r"path-to-experiment-log.xlsx"

DATA_SHEET_NAME = None   # e.g. "Sheet1" or leave as None for active sheet
SUMMARY_SHEET_NAME = "Min_Max_Summary"

HEADER_ROW = 1
NAME_COLUMN_HEADER = "run_id"   # change if your name column has a different header
START_COL = "AK"
END_COL = "BB"

# =========================================================
# STYLES
# =========================================================
RED_FILL = PatternFill(fill_type="solid", start_color="FFC7CE", end_color="FFC7CE")
GREEN_FILL = PatternFill(fill_type="solid", start_color="C6EFCE", end_color="C6EFCE")

# =========================================================
# HELPERS
# =========================================================
def get_header_map(ws, header_row=1):
    header_map = {}
    for cell in ws[header_row]:
        if cell.value is not None:
            header_map[str(cell.value).strip().lower()] = cell.column
    return header_map

def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

# =========================================================
# LOAD WORKBOOK
# =========================================================
file_path = Path(FILE_PATH)
wb = load_workbook(file_path)

ws = wb[DATA_SHEET_NAME] if DATA_SHEET_NAME else wb.active
header_map = get_header_map(ws, HEADER_ROW)

if NAME_COLUMN_HEADER.lower() not in header_map:
    raise ValueError(f"Could not find name column header: '{NAME_COLUMN_HEADER}'")

name_col_idx = header_map[NAME_COLUMN_HEADER.lower()]
start_col_idx = ws[START_COL + "1"].column
end_col_idx = ws[END_COL + "1"].column

# =========================================================
# PART 1: HIGHLIGHT LOWEST AND HIGHEST IN EACH ROW (AK:BB)
# =========================================================
for row_idx in range(HEADER_ROW + 1, ws.max_row + 1):
    numeric_cells = []

    for col_idx in range(start_col_idx, end_col_idx + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if is_number(cell.value):
            numeric_cells.append(cell)

    if not numeric_cells:
        continue

    values = [cell.value for cell in numeric_cells]
    row_min = min(values)
    row_max = max(values)

    for cell in numeric_cells:
        # If all values are the same, only color them green
        if row_min == row_max:
            if cell.value == row_max:
                cell.fill = GREEN_FILL
        else:
            if cell.value == row_min:
                cell.fill = RED_FILL
            if cell.value == row_max:
                cell.fill = GREEN_FILL

# =========================================================
# PART 2: CREATE/REPLACE SUMMARY SHEET
# Find which row and name had the min and max in each column
# =========================================================
if SUMMARY_SHEET_NAME in wb.sheetnames:
    del wb[SUMMARY_SHEET_NAME]

summary_ws = wb.create_sheet(SUMMARY_SHEET_NAME)

summary_headers = [
    "Column",
    "Minimum Value",
    "Min Row Number",
    "Min Name",
    "Maximum Value",
    "Max Row Number",
    "Max Name",
]
summary_ws.append(summary_headers)

for col_idx in range(start_col_idx, end_col_idx + 1):
    header_value = ws.cell(row=HEADER_ROW, column=col_idx).value

    numeric_entries = []
    for row_idx in range(HEADER_ROW + 1, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=col_idx).value
        if is_number(value):
            name_value = ws.cell(row=row_idx, column=name_col_idx).value
            numeric_entries.append((row_idx, name_value, value))

    if not numeric_entries:
        summary_ws.append([header_value, None, None, None, None, None, None])
        continue

    min_entry = min(numeric_entries, key=lambda x: x[2])  # (row, name, value)
    max_entry = max(numeric_entries, key=lambda x: x[2])

    summary_ws.append([
        header_value,
        min_entry[2],
        min_entry[0],
        min_entry[1],
        max_entry[2],
        max_entry[0],
        max_entry[1],
    ])

# =========================================================
# OPTIONAL: AUTO WIDTH
# =========================================================
for sheet in [summary_ws]:
    for col in sheet.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[col_letter].width = max_length + 2

# =========================================================
# SAVE
# =========================================================
wb.save(file_path)
print(f"Done. Workbook updated: {file_path}")
